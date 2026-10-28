"""Перенос журнала из Excel: tools/import_journal.py."""
import datetime
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from sz import actions  # noqa: E402
from sz.store import Store  # noqa: E402

from test_core import seed  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

HEADER = ['№', 'Объект', 'Номер СЗ', 'Номер СЗ в ИСУ КБ', 'Номер СЗ 1С', 'Краткое  содержание', 'Дата',
          'Ответственный', 'Кому']


def make_journal(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Журнал 2026'
    ws.append(['Журнал служебных записок за 2026 год'])
    ws.merge_cells('A1:I1')
    ws.append([])
    ws.append(HEADER)
    rows = [
        [1, 'Закупка', 'НТЦ-0001-26', 'ИСУ-1', '1С-1', 'Закупка ноутбуков', datetime.datetime(2026, 1, 15),
         'Сидоров Сидор Сидорович', 'Отдел кадров'],
        [2, 'Ремонт помещений', 'НТЦ-2-26', '', '', 'Ремонт кабинета 101', '05.02.2026 г.',
         'Сидоров С.С.', 'Петрову П.П.; Бухгалтерия'],
        [3, 'Командировки', 'НТЦ – 0003 – 26', '', '', 'Командировка в Москву', '10.02.26',
         'Сидоров', 'Начальнику отдела кадров Петрову П.П.\nПланово-экономический отдел'],
        [4, None, 'НТЦ-0004-26', '', '', 'Вторая командировка', 46075, 'Неизвестный Н.Н.', ''],
        [5, 'Закупка', 'без номера', '', '', 'Ошибка', '01.03.2026', '', ''],
        [],
        [6, 'Закупка', 'НТЦ-0006-26', '', '', 'Дата пустая', None, '', ''],
        [7, 'Закупка', 'НТЦ-0001-26', '', '', 'Повтор номера', '02.03.2026', '', ''],
        [8, 'Закупка', 'НТЦ-0050-26', '', '', 'Уже есть в базе', '03.03.2026', '', ''],
        [9, 'Закупка', 'НТЦ-0007-25', '', '', 'Прошлый год', '30.12.2025', '', 'Отдел кадров'],
    ]
    for r in rows:
        ws.append(r)
    ws.merge_cells('B6:B7')  # «Командировки» на две строки (4-я строка журнала без «Объекта»)
    wb.save(path)


@unittest.skipIf(openpyxl is None, 'openpyxl не установлен')
class ImportJournalTest(unittest.TestCase):
    def setUp(self):
        self.root = tempfile.mkdtemp()
        store = Store(self.root)
        store.open()
        seed(store)
        # Записка, которая уже есть в базе
        store.write(lambda s: s.run(
            "INSERT INTO memos (number, year, reg_number, memo_date, topic_name, addressee_department, "
            "addressee_position, addressee_name, body, executor_name) VALUES "
            "(50, 2026, 'НТЦ-0050-26', '2026-03-03', 'Т', 'О', 'Д', 'Ф', 'было', 'И')"))
        store.close()
        self.xlsx = os.path.join(self.root, 'журнал.xlsx')
        make_journal(self.xlsx)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def run_import(self):
        res = subprocess.run([sys.executable, os.path.join(ROOT, 'tools', 'import_journal.py'), self.xlsx,
                              self.root, '--yes'], stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
        out = res.stdout.decode('utf-8')
        self.assertEqual(res.returncode, 0, out)
        return out

    def test_import(self):
        out = self.run_import()
        self.assertIn('Будет внесено записок: 5', out)
        self.assertIn('Уже есть в базе и будут пропущены: 1', out)
        self.assertIn('непонятный номер СЗ «без номера»', out)
        self.assertIn('НТЦ-0006-26): непонятная дата', out)
        self.assertIn('номер НТЦ-0001-26 уже был в строке 4', out)

        store = Store(self.root)
        store.open()
        memos = {m['reg_number']: m for m in store.all('SELECT * FROM memos')}
        self.assertEqual(sorted(memos), ['НТЦ-0001-26', 'НТЦ-0002-26', 'НТЦ-0003-26', 'НТЦ-0004-26',
                                         'НТЦ-0007-25', 'НТЦ-0050-26'])
        m1 = memos['НТЦ-0001-26']
        self.assertEqual((m1['memo_date'], m1['topic_name'], m1['body'], m1['executor_name'], m1['executor_phone']),
                         ('2026-01-15', 'Закупка', 'Закупка ноутбуков', 'Сидоров Сидор Сидорович', '12-34'))
        self.assertEqual(m1['topic_id'], 1)  # тематика из справочника
        self.assertEqual(store.memo_addressees(m1['id'])[0]['head_name_dative'], 'Петрову П.П.')

        m2 = memos['НТЦ-0002-26']
        self.assertEqual(m2['memo_date'], '2026-02-05')
        self.assertEqual(m2['executor_id'], 1)  # «Сидоров С.С.» найден в справочнике
        to2 = store.memo_addressees(m2['id'])
        self.assertEqual([(a['addressee_id'], a['department']) for a in to2],
                         [(1, 'Отдел кадров'), (None, 'Бухгалтерия')])
        # В документе адресат, которого нет в справочнике, печатается как в журнале, без пустых строк
        from sz import docx
        blocks = [b.get('text') for b in docx.memo_blocks(m2, to2) if b.get('align') == 'right']
        self.assertEqual(blocks, ['Начальнику отдела кадров', 'Петрову П.П.', 'Бухгалтерия'])
        self.assertIsNotNone(m2['topic_id'])  # новая тематика добавлена в справочник
        self.assertTrue(store.one("SELECT * FROM topics WHERE name = 'Ремонт помещений'"))

        m3 = memos['НТЦ-0003-26']
        self.assertEqual((m3['memo_date'], m3['executor_id']), ('2026-02-10', 1))  # только фамилия
        self.assertEqual([a['department'] for a in store.memo_addressees(m3['id'])],
                         ['Отдел кадров', 'Планово-экономический отдел'])

        m4 = memos['НТЦ-0004-26']
        self.assertEqual(m4['topic_name'], 'Командировки')  # из объединённой ячейки
        self.assertEqual(m4['memo_date'], '2026-02-22')  # число Excel 46075
        self.assertEqual((m4['executor_id'], m4['executor_name']), (None, 'Неизвестный Н.Н.'))

        self.assertEqual(memos['НТЦ-0007-25']['year'], 2025)
        self.assertEqual(memos['НТЦ-0050-26']['body'], 'было')  # существующая не тронута

        # Найдутся поиском, нумерация продолжается после журнала
        self.assertEqual(len(actions.search_memos(store, 'ноутбук')), 1)
        self.assertEqual(actions.next_number(store, 2026), 51)
        self.assertEqual(actions.next_number(store, 2025), 8)
        store.close()

        # Повторный запуск ничего не добавляет
        out = self.run_import()
        self.assertIn('Будет внесено записок: 0', out)
        self.assertIn('Вносить нечего', out)
        backups = [f for f in os.listdir(os.path.join(self.root, 'data')) if f.startswith('memos-до-импорта')]
        self.assertEqual(len(backups), 1)


if __name__ == '__main__':
    unittest.main()
